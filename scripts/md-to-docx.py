#!/usr/bin/env python3
"""
Markdown 보고서 → Word(.docx) 변환기 (이 저장소 보고서용 최소 구현)

    python3 scripts/md-to-docx.py reports/보고서.md reports/보고서.docx

지원: 제목(#~####) · 문단 · 목록(-, 1.) · 표(| … |) · 코드블록(```) · 인용(>) · 구분선(---)
      · 인라인 **굵게** · `코드`
의존성: python-docx  (pip install python-docx)
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

KOREAN_FONT = "맑은 고딕"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x0B, 0x0B, 0x0B)
MUTED = RGBColor(0x6B, 0x6B, 0x72)


def set_font(run, name=KOREAN_FONT, size=None, bold=None, color=None, mono=False):
    """한글·영문 폰트를 함께 지정합니다(동아시아 폰트는 rPr/eastAsia 로 따로 넣어야 합니다)."""
    font = run.font
    font.name = MONO_FONT if mono else name
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    target = "eastAsia"
    rfonts.set(qn(f"w:{target}"), KOREAN_FONT if not mono else MONO_FONT)
    return run


def style_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = KOREAN_FONT
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), KOREAN_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    for level, size in ((1, 20), (2, 15), (3, 12.5), (4, 11)):
        st = doc.styles[f"Heading {level}"]
        st.font.name = KOREAN_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), KOREAN_FONT)
        st.paragraph_format.space_before = Pt(14 if level > 1 else 0)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text, size=None, base_bold=False, color=None):
    """**굵게** · `코드` 를 서식으로 바꿔 문단에 넣습니다."""
    text = text.replace("&#124;", "|")
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            set_font(paragraph.add_run(part[2:-2]), size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_font(run, size=(size or 10) - 0.5, mono=True, color=RGBColor(0x8A, 0x3D, 0x00))
            run.font.highlight_color = None
        else:
            set_font(paragraph.add_run(part), size=size, bold=base_bold or None, color=color)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, "F4F4F5")
    cell.paragraphs[0].text = ""
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_font(p.add_run(line), size=9, mono=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].text = ""
        shade(cell, "EFEFEF")
        add_inline(cell.paragraphs[0], text, size=9, base_bold=True)
    for row in body:
        cells = table.add_row().cells
        for i, text in enumerate(row[: len(header)]):
            cells[i].paragraphs[0].text = ""
            add_inline(cells[i].paragraphs[0], text, size=9)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    cells, buf, in_code = [], "", False
    for ch in line:
        if ch == "`":
            in_code = not in_code
        if ch == "|" and not in_code:
            cells.append(buf.strip())
            buf = ""
        else:
            buf += ch
    cells.append(buf.strip())
    return cells


def convert(md_path, docx_path):
    raw = open(md_path, encoding="utf-8").read().splitlines()
    doc = Document()
    style_document(doc)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2)

    i = 0
    while i < len(raw):
        line = raw[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(raw) and not raw[i].strip().startswith("```"):
                block.append(raw[i])
                i += 1
            add_code_block(doc, block)
            i += 1
            continue

        if stripped in ("---", "***"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            set_font(p.add_run("─" * 46), size=8, color=RGBColor(0xC8, 0xC8, 0xCE))
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(raw) and re.match(r"^\|[\s:\-|]+\|?$", raw[i + 1].strip()):
            rows = [split_row(stripped)]
            i += 2
            while i < len(raw) and raw[i].strip().startswith("|"):
                rows.append(split_row(raw[i]))
                i += 1
            add_table(doc, rows)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1 and not doc.paragraphs:
                p = doc.add_paragraph()
                add_inline(p, text, size=22, base_bold=True)
                p.paragraph_format.space_after = Pt(10)
            else:
                heading = doc.add_heading(level=min(level, 4))
                add_inline(heading, text, size=None, base_bold=True)
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^[-*]\s+", "", stripped))
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, stripped.lstrip("> ").strip(), color=MUTED)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(docx_path)
    print(f"✔ {md_path} → {docx_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
